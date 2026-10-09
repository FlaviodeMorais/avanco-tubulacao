#!/usr/bin/env python3
"""De-para por spool: SOP/SUBSOP/STH antigos -> novos, a partir do CONTROLE GERAL DE STH-HC².

Uso:
  python3 scripts/gerar_de_para_spool.py "CONTROLE GERAL DE STH-HC²-24-09.xlsx" ON-SITE_Spools.xlsx [de_para_spool.json]

Aba "CONTROLE BASE" (cabeçalho na linha 4): 'Nº do Spool' (iso-spool), 'SOP  NOVO', 'SUBSOP NOVO',
'STH NOVO' (SOP-STH, ex. STH-U-4710-0009-0034 -> STH do P6 = STH-U-4710-0034), 'LINHA'.
Liga cada spool do ON-SITE por isométrico-spool (fallback: linha + nº do spool).
Saída: {"iso|spool": [sop_novo, subsop_novo, sth_p6, sth_novo]} e a lista de cancelados.
"""
import json, re, sys
from collections import Counter
import openpyxl


def sth_p6(s):
    m = re.match(r'(STH-[A-Z]-\d{4})-\d{4}-(\d{4})$', str(s or '').strip())
    return f'{m.group(1)}-{m.group(2)}' if m else None


def main(ctl_xlsx, onsite_xlsx, dst):
    ws = openpyxl.load_workbook(ctl_xlsx, read_only=True, data_only=True)['CONTROLE BASE']
    rows = list(ws.iter_rows(values_only=True))
    ix = {str(n).strip(): i for i, n in enumerate(rows[3]) if n}
    ctl, por_linha = {}, {}
    for r in rows[4:]:
        if not r[ix['Nº do Spool']]:
            continue
        num = str(r[ix['Nº do Spool']]).strip()
        reg = (r[ix['SOP  NOVO']], r[ix['SUBSOP NOVO']], r[ix['STH NOVO']])
        ctl[num] = reg
        por_linha[(str(r[ix['LINHA']] or '').strip(), num[-3:])] = reg

    ws = openpyxl.load_workbook(onsite_xlsx, read_only=True, data_only=True)['Mapa de Spools']
    it = ws.iter_rows(min_row=6, values_only=True)
    cab = {str(c or '').strip(): i for i, c in enumerate(next(it))}
    out, cancel, stat = {}, [], Counter()
    for r in it:
        if not r[cab['Unidade']] or r[cab['Peso']] is None:
            continue
        iso, sp0 = str(r[cab['Isométrico']]).strip(), str(r[cab['Spool']]).strip()
        sp = sp0.zfill(3)
        reg = ctl.get(f'{iso}-{sp}') or ctl.get(f'{iso}-{sp0}')
        via = 'iso-spool'
        if reg is None:
            reg = por_linha.get((str(r[cab['Linha']] or '').strip(), sp)); via = 'linha+spool'
        if reg is None:
            stat['sem correspondência'] += 1; continue
        stat[via] += 1
        sop, ssop, sth = reg
        if 'CANCELAD' in str(sth).upper() or 'CANCELAD' in str(sop).upper():
            cancel.append(f'{r[cab["Isométrico"]]}|{r[cab["Spool"]]}')
            stat['cancelados'] += 1
        key = f'{r[cab["Isométrico"]]}|{r[cab["Spool"]]}'
        out[key] = [sop, ssop, sth_p6(sth) or None, sth]
    json.dump({'fonte': 'CONTROLE GERAL DE STH-HC2 (24-09) · aba CONTROLE BASE', 'spools': out, 'cancelados': cancel},
              open(dst, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    print(dict(stat), '· STH P6:', len({v[2] for v in out.values() if v[2]}),
          '· SOP novos:', len({v[0] for v in out.values() if v[0]}))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else 'de_para_spool.json')
